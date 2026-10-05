"""Offline tests for the pipeline. No network, no GLM.

Run:  .venv/bin/python -m unittest discover -s tests -v
"""
import copy
import json
import os
import shutil
import tempfile
import unittest

import openpyxl

from src import consolidate
from src.excel_writer import write_excel
from src.glm_structurer import StructureError, parse_reply
from src.scraper import is_match, make_notice_id, parse_listing
from src.state import State
from src.validator import firm_key, normalize_category, normalize_date, validate

PAGE_URL = "https://www.mod.gov.in/dod/en/oms_orders_Archive?page=3"

LISTING_HTML = """
<html><body><table>
<thead><tr><th>S.N</th><th>Title</th><th>Date</th><th>Download</th></tr></thead>
<tbody>
<tr>
  <td class="views-field views-field-counter">1</td>
  <td class="views-field views-field-title">Details of firms debarred/put on hold/suspended
      etc. from doing business with MoD-reg.</td>
  <td class="views-field views-field-field-start-date"><time datetime="2025-09-01T12:00:00Z">2025-09-01</time></td>
  <td class="views-field views-field-nothing"><a href="/dod/sites/default/files/details-of-firms.pdf">Download</a></td>
</tr>
<tr>
  <td class="views-field views-field-counter">2</td>
  <td class="views-field views-field-title">Posting/transfer at Director level</td>
  <td class="views-field views-field-field-start-date"><time>2025-08-01</time></td>
  <td class="views-field views-field-nothing"><a href="/dod/sites/default/files/posting.pdf">Download</a></td>
</tr>
<tr>
  <td class="views-field views-field-counter">3</td>
  <td class="views-field views-field-title">Order without a file</td>
  <td class="views-field views-field-field-start-date"><time>2025-07-01</time></td>
  <td class="views-field views-field-nothing"></td>
</tr>
</tbody></table></body></html>
"""

GLM_REPLY = """```json
{"reference_no": "C-31013/1/2025-D(Vig.)-2", "as_on_date": "01.09.2025",
 "sections": [
  {"category": "Debarred Firms", "stated_total": 2, "firms": [
    {"s_no": 1, "firm_name": "M/s Singapore Technologies Kinetics Ltd. (STK)", "location": null,
     "mod_division": "D(Vig)", "order_ref": "MoD ID", "order_date": "11.04.2012"},
    {"s_no": 2, "firm_name": "M/s R. K. Machine Tools Ltd.", "location": "Ludhiana",
     "mod_division": "NDCD", "order_ref": "Order No. 22(4)/2025/NDCD/DDP", "order_date": "01.04.2025"}]},
  {"category": "Put on hold/Suspended firms", "stated_total": 2, "firms": [
    {"s_no": 1, "firm_name": "M/s IDS", "location": "Tunisia", "mod_division": "Air",
     "order_ref": null, "order_date": "03.07.2014"}]}
 ]}
```"""


class TestKeywordMatch(unittest.TestCase):
    def test_matches_all_forms(self):
        for title in ["Debarment of firm X", "Firms DEBARRED from business",
                      "debar order", "Details of firms debarred/put on hold"]:
            self.assertTrue(is_match(title, "debar"), title)

    def test_ignores_unrelated_titles(self):
        self.assertFalse(is_match("Posting/transfer at Director level", "debar"))

    def test_handles_whitespace_and_newlines(self):
        self.assertTrue(is_match("Details of firms\n\t  Debarred", "debar"))

    def test_scraper_does_not_use_regex(self):
        """User requirement: plain string matching only."""
        with open(os.path.join(os.path.dirname(__file__), "..", "src", "scraper.py")) as f:
            source = f.read()
        self.assertNotIn("import re", source)


class TestParseListing(unittest.TestCase):
    def setUp(self):
        self.rows = parse_listing(LISTING_HTML, PAGE_URL)

    def test_reads_every_row(self):
        self.assertEqual(len(self.rows), 3)

    def test_fields(self):
        r = self.rows[0]
        self.assertEqual(r["listed_date"], "2025-09-01")
        self.assertEqual(r["pdf_url"], "https://www.mod.gov.in/dod/sites/default/files/details-of-firms.pdf")
        self.assertEqual(r["title"], "Details of firms debarred/put on hold/suspended etc. "
                                     "from doing business with MoD-reg.")

    def test_missing_link_is_none(self):
        self.assertIsNone(self.rows[2]["pdf_url"])

    def test_no_table_returns_empty(self):
        self.assertEqual(parse_listing("<html><body>The specified URL cannot be found.</body></html>", PAGE_URL), [])

    def test_notice_id(self):
        self.assertEqual(make_notice_id("2025-09-01", self.rows[0]["pdf_url"]), "2025-09-01_details-of-firms")


class TestNormalize(unittest.TestCase):
    def test_dates(self):
        self.assertEqual(normalize_date("01.09.2025"), ("2025-09-01", True))
        self.assertEqual(normalize_date("1/9/25"), ("2025-09-01", True))
        self.assertEqual(normalize_date("01-09-2025"), ("2025-09-01", True))
        self.assertEqual(normalize_date("2025-09-01"), ("2025-09-01", True))
        self.assertEqual(normalize_date("31.13.2025"), ("31.13.2025", False))
        self.assertEqual(normalize_date("unreadable"), ("unreadable", False))
        self.assertEqual(normalize_date(None), (None, False))

    def test_firm_key_matches_ocr_variants(self):
        expected = "r k machine tools ltd"
        for name in ["M/s R. K. Machine Tools Ltd.", "Ms R.K. Machine Tools Ltd",
                     "M/S  R K Machine Tools Ltd.,", "m/s. R. K. Machine Tools Ltd"]:
            self.assertEqual(firm_key(name), expected, name)

    def test_category(self):
        self.assertEqual(normalize_category("Debarred Firms - Total 6"), "Debarred")
        self.assertEqual(normalize_category("Put on hold/Suspended firms"), "Put on hold/Suspended")
        self.assertEqual(normalize_category("Suspended"), "Put on hold/Suspended")


class TestParseReply(unittest.TestCase):
    def test_strips_code_fence(self):
        data = parse_reply(GLM_REPLY)
        self.assertEqual(len(data["sections"]), 2)

    def test_rejects_non_json(self):
        with self.assertRaises(StructureError):
            parse_reply("Sorry, I cannot read this document.")

    def test_rejects_missing_sections(self):
        with self.assertRaises(StructureError):
            parse_reply('{"reference_no": "x"}')


class TempDirTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.cfg = {"paths": {"json": self.tmp}}

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def doc(self, notice_id, listed_date):
        return {"notice_id": notice_id, "listed_date": listed_date, "title": "t",
                "sha256": "abc", "pdf_url": f"https://example.test/{notice_id}.pdf"}


class TestValidate(TempDirTest):
    def test_total_mismatch_flags_section(self):
        clean = validate(parse_reply(GLM_REPLY), self.doc("n1", "2025-09-01"), self.cfg)
        self.assertEqual(clean["as_on_date"], "2025-09-01")
        self.assertIn("Put on hold/Suspended: PDF says Total 2, extracted 1", clean["issues"])
        debarred, hold = clean["sections"]
        self.assertFalse(any(f["needs_review"] for f in debarred["firms"]))
        self.assertTrue(hold["firms"][0]["needs_review"])
        self.assertIn("section total mismatch", hold["firms"][0]["review_reason"])
        self.assertIn("no order ref", hold["firms"][0]["review_reason"])

    def test_writes_clean_json(self):
        validate(parse_reply(GLM_REPLY), self.doc("n1", "2025-09-01"), self.cfg)
        with open(os.path.join(self.tmp, "n1.clean.json")) as f:
            self.assertEqual(json.load(f)["notice_id"], "n1")

    def test_bad_as_on_date_falls_back_to_listed_date(self):
        ext = parse_reply(GLM_REPLY)
        ext["as_on_date"] = None
        clean = validate(ext, self.doc("n1", "2025-09-01"), self.cfg)
        self.assertEqual(clean["as_on_date"], "2025-09-01")
        self.assertTrue(any("as_on_date" in i for i in clean["issues"]))


class TestConsolidateAndExcel(TempDirTest):
    def setUp(self):
        super().setUp()
        new = parse_reply(GLM_REPLY)
        old = copy.deepcopy(new)
        old["as_on_date"] = "03.02.2023"
        old["sections"][0]["firms"][1]["firm_name"] = "Ms R.K. Machine Tools Ltd"   # same firm, OCR variant
        old["sections"][1]["firms"][0]["firm_name"] = "M/s Experts Systems"         # later removed
        old["sections"][0]["firms"][0]["firm_name"] = "M/s IDS"                      # later moved category
        new["sections"][0]["firms"][0]["firm_name"] = "M/s Singapore Technologies Kinetics Ltd. (STK)"
        validate(old, self.doc("2023-02-10_a", "2023-02-10"), self.cfg)
        validate(new, self.doc("2025-09-01_b", "2025-09-01"), self.cfg)
        self.notices = consolidate.load_clean(self.tmp)
        self.records = consolidate.to_records(self.notices)

    def test_records_one_per_firm(self):
        self.assertEqual(len(self.records), 6)

    def test_current_status_is_latest_snapshot(self):
        current = consolidate.current_status(self.notices, self.records)
        self.assertEqual({r["notice_id"] for r in current}, {"2025-09-01_b"})

    def test_changes(self):
        found = {(c["change"], c["firm_name"]) for c in consolidate.changes(self.notices, self.records)}
        self.assertIn(("Removed", "M/s Experts Systems"), found)
        self.assertIn(("Added", "M/s Singapore Technologies Kinetics Ltd. (STK)"), found)
        self.assertIn(("Category changed", "M/s IDS"), found)
        # OCR variant of the same firm must NOT show up as a change
        self.assertFalse(any("Machine Tools" in name for _, name in found))

    def test_excel_has_four_sheets(self):
        path = os.path.join(self.tmp, "out.xlsx")
        write_excel(path, consolidate.current_status(self.notices, self.records), self.records,
                    consolidate.changes(self.notices, self.records),
                    [{"notice_id": "2025-09-01_b", "status": "success", "firms": 3}])
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Current Status", "All Records", "Changes", "Run Log"])
        self.assertEqual(wb["All Records"].max_row - 1, 6)
        link = wb["All Records"]["K2"]
        self.assertEqual(link.value, "Open PDF")
        self.assertTrue(link.hyperlink.target.startswith("https://"))


class TestState(TempDirTest):
    def test_done_and_sha_change(self):
        path = os.path.join(self.tmp, "state", "processed.json")
        s = State(path)
        url = "https://example.test/a.pdf"
        self.assertFalse(s.is_done(url))
        s.mark(url, "done", sha256="aaa")
        s2 = State(path)  # reloaded from disk
        self.assertTrue(s2.is_done(url))
        self.assertFalse(s2.sha_changed(url, "aaa"))
        self.assertTrue(s2.sha_changed(url, "bbb"))
        self.assertFalse(s2.sha_changed("https://example.test/unknown.pdf", "bbb"))

    def test_failed_is_not_done(self):
        s = State(os.path.join(self.tmp, "p.json"))
        s.mark("u", "failed", error="timeout")
        self.assertFalse(s.is_done("u"))


if __name__ == "__main__":
    unittest.main()
