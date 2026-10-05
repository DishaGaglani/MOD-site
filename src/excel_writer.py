"""S8 Export: rebuild output/debarment_list.xlsx from the consolidated data."""
import os

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")

CURRENT_COLS = {
    "category": "Category", "s_no": "S.No", "firm_name": "Firm Name", "location": "Location",
    "mod_division": "MoD Division", "order_ref": "Order No.", "order_date": "Order Date",
    "as_on_date": "As-On Date", "pdf_url": "PDF Link", "needs_review": "Needs Review",
}
RECORD_COLS = {
    "listed_date": "Notice Date", "as_on_date": "As-On Date", "reference_no": "Reference No.",
    "category": "Category", "s_no": "S.No", "firm_name": "Firm Name", "location": "Location",
    "mod_division": "MoD Division", "order_ref": "Order No.", "order_date": "Order Date",
    "pdf_url": "PDF Link", "needs_review": "Needs Review", "review_reason": "Review Reason",
}
CHANGE_COLS = {
    "as_on_date": "As-On Date", "previous_as_on_date": "Compared With", "firm_name": "Firm Name",
    "change": "Change", "from": "From", "to": "To", "pdf_url": "PDF Link",
}
LOG_COLS = {
    "notice_id": "Notice", "listed_date": "Notice Date", "status": "Status",
    "firms": "Firms Extracted", "needs_review": "Rows Needing Review",
    "issues": "Issues / Error", "pdf_url": "PDF Link",
}


def _frame(rows, cols):
    return pd.DataFrame(rows, columns=list(cols)).rename(columns=cols)


def write_excel(path, current, records, changes, run_log):
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    sheets = {
        "Current Status": _frame(current, CURRENT_COLS),
        "All Records": _frame(records, RECORD_COLS),
        "Changes": _frame(changes, CHANGE_COLS),
        "Run Log": _frame(run_log, LOG_COLS),
    }
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            _format(xw.sheets[name], df)


def _format(ws, df):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    headers = list(df.columns)
    for idx, col in enumerate(headers, start=1):
        values = [str(v) for v in df[col].tolist() if v is not None]
        width = max([len(col)] + [len(v) for v in values]) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width, 10), 60)

    link_col = headers.index("PDF Link") + 1 if "PDF Link" in headers else None
    review_col = headers.index("Needs Review") + 1 if "Needs Review" in headers else None
    for row in ws.iter_rows(min_row=2):
        if link_col:
            c = row[link_col - 1]
            if c.value:
                c.hyperlink = c.value
                c.value = "Open PDF"
                c.font = Font(color="0563C1", underline="single")
        if review_col and row[review_col - 1].value is True:
            for c in row:
                if c.column != link_col:
                    c.fill = REVIEW_FILL
